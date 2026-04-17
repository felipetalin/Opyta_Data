from __future__ import annotations

from io import BytesIO

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


GROUP_TEMPLATE_SPECS: dict[str, dict[str, object]] = {
    "Ictiofauna": {
        "filename": "modelo_oficial_ictiofauna_v1_0.xlsx",
        "sheet_name": "Resultados_Ictiofauna",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Quantidade",
            "Biomassa",
            "Artefato_Pesca",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Astyanax sp.", 12, 1.45, "Rede de espera"],
    },
    "Bentos": {
        "filename": "modelo_oficial_bentos_v1_0.xlsx",
        "sheet_name": "Resultados_Zoobentos",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Quantidade",
            "BMWP_Score",
            "Ordem_Taxonomica",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Chironomidae sp.", 34, 5, "Diptera"],
    },
    "Fitoplâncton": {
        "filename": "modelo_oficial_fitoplancton_v1_0.xlsx",
        "sheet_name": "Resultados_Fitoplancton",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Densidade",
            "Biovolume",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Microcystis sp.", 1520, 0.84],
    },
    "Zooplâncton": {
        "filename": "modelo_oficial_zooplancton_v1_0.xlsx",
        "sheet_name": "Resultados_Zooplancton",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Quantidade",
            "Volume_Filtrado",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Brachionus sp.", 87, 120],
    },
    "Meio Físico": {
        "filename": "modelo_oficial_meio_fisico_v1_0.xlsx",
        "sheet_name": "Resultados_Meio_Fisico",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Matriz",
            "Nome_Parametro",
            "Valor_Medido",
            "Unidade_Medida",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Água Superficial", "pH", 7.2, "pH"],
    },
    "Avifauna": {
        "filename": "modelo_oficial_avifauna_v1_0.xlsx",
        "sheet_name": "Resultados_Avifauna",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Quantidade",
            "Metodo_Amostragem",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Turdus rufiventris", 3, "Ponto fixo"],
    },
    "Herpetofauna": {
        "filename": "modelo_oficial_herpetofauna_v1_0.xlsx",
        "sheet_name": "Resultados_Herpetofauna",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Quantidade",
            "Metodo_Amostragem",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Boana punctata", 4, "Busca ativa"],
    },
    "Mastofauna": {
        "filename": "modelo_oficial_mastofauna_v1_0.xlsx",
        "sheet_name": "Resultados_Mastofauna",
        "headers": [
            "Campanha",
            "Ponto",
            "Data_Coleta",
            "Latitude",
            "Longitude",
            "Nome_Cientifico",
            "Quantidade",
            "Metodo_Amostragem",
        ],
        "sample": ["2026-01", "P01", "2026-01-15", -2.5301, -44.3021, "Didelphis marsupialis", 1, "Armadilha fotográfica"],
    },
}


def _apply_header_style(cell) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")


@st.cache_data(show_spinner=False)
def build_group_template_bytes(group: str) -> bytes:
    spec = GROUP_TEMPLATE_SPECS[group]
    workbook = Workbook()

    ws_instrucoes = workbook.active
    ws_instrucoes.title = "Leia-me"
    ws_instrucoes["A1"] = "Modelo oficial Opyta"
    ws_instrucoes["A2"] = f"Grupo: {group}"
    ws_instrucoes["A4"] = "Regras rápidas"
    ws_instrucoes["A5"] = "1. Não altere os nomes das abas."
    ws_instrucoes["A6"] = "2. Mantenha os nomes das colunas exatamente como no modelo."
    ws_instrucoes["A7"] = "3. Use uma linha por registro coletado."
    ws_instrucoes["A8"] = "4. Preencha latitude/longitude em decimal."
    ws_instrucoes["A9"] = "5. Revise o arquivo antes de subir na importação."

    ws_capa = workbook.create_sheet("Capa_Projeto")
    ws_capa.append(["Codigo_Opyta", "Nome_Projeto", "Responsavel", "Versao_Modelo"])
    ws_capa.append(["OP-001", "Projeto Exemplo", "Equipe Opyta", "v1.0"])

    ws_pontos = workbook.create_sheet("Pontos_e_Campanhas")
    ws_pontos.append(["Campanha", "Ponto", "Latitude", "Longitude", "Data_Coleta"])
    ws_pontos.append(["2026-01", "P01", -2.5301, -44.3021, "2026-01-15"])
    ws_pontos.append(["2026-01", "P02", -2.5312, -44.3034, "2026-01-16"])

    ws_resultados = workbook.create_sheet(str(spec["sheet_name"]))
    headers = list(spec["headers"])
    ws_resultados.append(headers)
    ws_resultados.append(list(spec["sample"]))

    for sheet in [ws_capa, ws_pontos, ws_resultados]:
        for cell in sheet[1]:
            _apply_header_style(cell)
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 4, 28)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def get_group_template_filename(group: str) -> str:
    return str(GROUP_TEMPLATE_SPECS[group]["filename"])